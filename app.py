
import streamlit as st
import pandas as pd
import numpy as np
from datetime import date, timedelta, datetime
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import re
import os

st.set_page_config(page_title="Break Compliance Report", page_icon="📊", layout="centered")

st.markdown("# 📊 Break Compliance Report")
st.markdown("Upload attendance CSV → Get formatted report instantly!")
st.markdown("---")

# --- Load Department Mapping from Excel file ---
@st.cache_data
def load_dept_map():
    """Load department mapping from the HC Excel file in the same directory."""
    dept_map = {}
    # Look for any HC Excel file in the current directory
    hc_file = None
    for f in os.listdir('.'):
        if f.endswith('.xlsx') and 'HC' in f.upper() and 'DEPT' in f.upper():
            hc_file = f
            break
    
    if hc_file:
        try:
            hc_df = pd.read_excel(hc_file)
            # Find the Psoft ID and Department columns
            id_col = None
            dept_col = None
            for col in hc_df.columns:
                if 'psoft' in col.lower() or 'id' in col.lower():
                    id_col = col
                elif 'dept' in col.lower() or 'department' in col.lower():
                    dept_col = col
            
            if id_col and dept_col:
                for _, row in hc_df.iterrows():
                    if pd.notna(row[id_col]) and pd.notna(row[dept_col]):
                        dept_map[int(row[id_col])] = str(row[dept_col])
        except Exception as e:
            st.warning(f"Could not load HC file: {e}")
    
    return dept_map

DEPT_MAP = load_dept_map()

if DEPT_MAP:
    st.sidebar.success(f"✅ Loaded {len(DEPT_MAP)} employees from HC file")
else:
    st.sidebar.warning("⚠️ No HC file found. Place your 'HC ... DEPT.xlsx' file in the app directory.")

# --- Colors (matching desktop report) ---
SQUID_INK = '232F3E'
TEAL = '00BCD4'
CORAL = 'FF6B6B'
SUNSET = 'FFA726'
CHARCOAL = '424242'
SNOW = 'FAFAFA'
ICE_BLUE = 'E0F7FA'
LIGHT_CORAL = 'FFEBEE'
LIGHT_SUNSET = 'FFF3E0'
WHITE = 'FFFFFF'
DARK_TEXT = '212121'
REPEAT_RED = 'D32F2F'
REPEAT_BG = 'FFCDD2'

# --- Helper: Extract date from filename ---
def extract_date_from_filename(filename):
    match = re.search(r'(\d{8})\d{4}-\d{12}', filename)
    if match:
        date_str = match.group(1)
        dt = datetime.strptime(date_str, '%Y%m%d')
        return dt.strftime('%A, %d %B %Y')
    
    match = re.search(r'(\d{4})(\d{2})(\d{2})', filename)
    if match:
        year, month, day = match.groups()
        dt = datetime(int(year), int(month), int(day))
        return dt.strftime('%A, %d %B %Y')
    
    match = re.search(r'(\d{4})-(\d{2})-(\d{2})', filename)
    if match:
        year, month, day = match.groups()
        dt = datetime(int(year), int(month), int(day))
        return dt.strftime('%A, %d %B %Y')
    
    yesterday = date.today() - timedelta(days=1)
    return yesterday.strftime('%A, %d %B %Y')

# --- Helper: Extract short date for history ---
def extract_short_date(filename):
    match = re.search(r'(\d{8})\d{4}-\d{12}', filename)
    if match:
        date_str = match.group(1)
        dt = datetime.strptime(date_str, '%Y%m%d')
        return f"{dt.month}/{dt.day}/{dt.year}"
    
    match = re.search(r'(\d{4})(\d{2})(\d{2})', filename)
    if match:
        year, month, day = match.groups()
        dt = datetime(int(year), int(month), int(day))
        return f"{dt.month}/{dt.day}/{dt.year}"
    
    match = re.search(r'(\d{4})-(\d{2})-(\d{2})', filename)
    if match:
        year, month, day = match.groups()
        dt = datetime(int(year), int(month), int(day))
        return f"{dt.month}/{dt.day}/{dt.year}"
    
    yesterday = date.today() - timedelta(days=1)
    return f"{yesterday.month}/{yesterday.day}/{yesterday.year}"

# --- Helper: Auto-detect shift from filename ---
def detect_shift_from_filename(filename):
    match = re.search(r'\d{8}(\d{4})-\d{12}', filename)
    if match:
        start_time = match.group(1)
        hour = int(start_time[:2])
        if hour >= 14:
            return 1  # Night Shift
        else:
            return 0  # Day Shift
    return 0

# --- Upload History ---
st.markdown("### 📋 Upload History (optional)")
history_file = st.file_uploader("Upload previous history.csv for repeat tracking", type=['csv'], key="history", help="Max 200MB per file")

history_df = pd.DataFrame(columns=['Employee ID', 'Employee Name', 'Date', 'Shift', 'Flag'])
if history_file:
    history_df = pd.read_csv(history_file)
    st.success(f"History loaded: {len(history_df)} past records")

# --- Upload Attendance CSV ---
st.markdown("### 📊 Upload Attendance CSV")
uploaded_file = st.file_uploader("Upload your attendance CSV file", type=['csv'], key="attendance", help="Max 200MB per file")

if uploaded_file:
    df = pd.read_csv(uploaded_file)
    filename = uploaded_file.name
    
    # Extract date from filename
    report_date = extract_date_from_filename(filename)
    short_date = extract_short_date(filename)
    
    # Auto-detect shift with override option
    default_shift = detect_shift_from_filename(filename)
    shift_type = st.selectbox("Select Shift:", ["Day Shift", "Night Shift"], index=default_shift)
    
    st.info(f"📅 Date: **{report_date}** | 🔄 Shift: **{shift_type}**")
    
    # --- Detect format ---
    if 'Punch Time' in df.columns:
        st.info("Detected: **FCLM Raw Punch Data** — calculating breaks from timestamps...")
        
        df['Punch Time'] = pd.to_datetime(df['Punch Time'])
        
        results = []
        for emp_id, group in df.groupby('Employee ID'):
            punches = group.sort_values('Punch Time')
            emp_name = punches['Employee Name'].iloc[0]
            
            if len(punches) <= 2:
                continue
            
            punch_times = punches['Punch Time'].tolist()
            total_break = 0
            
            for i in range(1, len(punch_times) - 1, 2):
                break_mins = (punch_times[i+1] - punch_times[i]).total_seconds() / 60
                if 0 < break_mins < 120:
                    total_break += break_mins
            
            if total_break == 0 and len(punch_times) >= 4:
                break_mins = (punch_times[2] - punch_times[1]).total_seconds() / 60
                if 0 < break_mins < 120:
                    total_break = break_mins
            
            if total_break > 0:
                results.append({
                    'Employee ID': emp_id,
                    'Employee Name': emp_name,
                    'Break (min)': round(total_break)
                })
        
        processed_df = pd.DataFrame(results)
    
    elif 'First Half Duration' in df.columns or 'Total Duration' in df.columns:
        st.info("Detected: **Dashboard Export** — reading break totals...")
        
        if 'First Half Duration' in df.columns and 'Second Half Duration' in df.columns:
            df['Break (min)'] = df['First Half Duration'].fillna(0) + df['Second Half Duration'].fillna(0)
        elif 'Total Duration' in df.columns:
            df['Break (min)'] = df['Total Duration'].fillna(0)
        
        processed_df = df[['Employee ID', 'Employee Name', 'Break (min)']].copy()
        processed_df = processed_df[processed_df['Break (min)'] > 0]
    
    else:
        st.error("❌ Unrecognized CSV format. Please upload FCLM Raw Punch Data or Dashboard Export.")
        st.stop()
    
    if len(processed_df) == 0:
        st.warning("No break data found to process.")
        st.stop()
    
    st.success(f"Processed **{len(processed_df)}** employees!")
    
    # --- Add Department ---
    processed_df['Department'] = processed_df['Employee ID'].map(DEPT_MAP).fillna('Unknown')
    
    # --- Flag Excess and Less ---
    excess_df = processed_df[processed_df['Break (min)'] >= 65].sort_values('Break (min)', ascending=False).reset_index(drop=True)
    less_df = processed_df[processed_df['Break (min)'] <= 55].sort_values('Break (min)', ascending=True).reset_index(drop=True)
    
    # --- Count repeats from history ---
    def count_repeats(emp_id, flag):
        if history_df.empty:
            return 0
        matches = history_df[(history_df['Employee ID'] == emp_id) & (history_df['Flag'] == flag)]
        return len(matches)
    
    excess_df['Repeat'] = excess_df['Employee ID'].apply(lambda x: count_repeats(x, 'Excess'))
    less_df['Repeat'] = less_df['Employee ID'].apply(lambda x: count_repeats(x, 'Less'))
    
    # Format repeat column with ⚠️ triangle
    excess_df['Repeat'] = excess_df['Repeat'].apply(lambda x: f"⚠️ {x}x" if x > 0 else "")
    less_df['Repeat'] = less_df['Repeat'].apply(lambda x: f"⚠️ {x}x" if x > 0 else "")
    
    # --- Display Metrics ---
    total = len(excess_df) + len(less_df)
    col1, col2, col3 = st.columns(3)
    col1.metric("Total Exceptions", total)
    col2.metric("Excess (≥65 min)", len(excess_df))
    col3.metric("Less (≤55 min)", len(less_df))
    
    # --- Display Tables ---
    if len(excess_df) > 0:
        st.markdown("### 🔴 Excess Break (≥65 min)")
        st.dataframe(excess_df[['Employee ID', 'Employee Name', 'Department', 'Break (min)', 'Repeat']], use_container_width=True)
    
    if len(less_df) > 0:
        st.markdown("### 🟡 Less Break (≤55 min)")
        st.dataframe(less_df[['Employee ID', 'Employee Name', 'Department', 'Break (min)', 'Repeat']], use_container_width=True)
    
    # --- Build Updated History ---
    new_history_records = []
    for _, row in excess_df.iterrows():
        new_history_records.append({
            'Employee ID': row['Employee ID'],
            'Employee Name': row['Employee Name'],
            'Date': short_date,
            'Shift': shift_type,
            'Flag': 'Excess'
        })
    for _, row in less_df.iterrows():
        new_history_records.append({
            'Employee ID': row['Employee ID'],
            'Employee Name': row['Employee Name'],
            'Date': short_date,
            'Shift': shift_type,
            'Flag': 'Less'
        })
    
    new_history_df = pd.DataFrame(new_history_records)
    updated_history = pd.concat([history_df, new_history_df], ignore_index=True)
    
    # --- Generate Excel Report (Matching Desktop Style) ---
    def generate_excel(excess, less, report_date, shift_type):
        wb = Workbook()
        ws = wb.active
        ws.title = "Break Compliance"
        ws.sheet_view.showGridLines = False
        
        # Column widths
        ws.column_dimensions['A'].width = 14
        ws.column_dimensions['B'].width = 28
        ws.column_dimensions['C'].width = 14
        ws.column_dimensions['D'].width = 8
        ws.column_dimensions['E'].width = 10
        
        # === HEADER BANNER ===
        for c in range(1, 6):
            ws.cell(row=1, column=c).fill = PatternFill(start_color=SQUID_INK, end_color=SQUID_INK, fill_type='solid')
        ws.row_dimensions[1].height = 10
        
        ws.row_dimensions[2].height = 35
        ws.merge_cells('A2:E2')
        ws['A2'] = "BREAK COMPLIANCE REPORT"
        ws['A2'].font = Font(name='Calibri', size=16, bold=True, color=TEAL)
        ws['A2'].fill = PatternFill(start_color=SQUID_INK, end_color=SQUID_INK, fill_type='solid')
        ws['A2'].alignment = Alignment(vertical='center', horizontal='center')
        for c in range(1, 6):
            ws.cell(row=2, column=c).fill = PatternFill(start_color=SQUID_INK, end_color=SQUID_INK, fill_type='solid')
        
        # Teal accent line
        ws.row_dimensions[3].height = 4
        for c in range(1, 6):
            ws.cell(row=3, column=c).fill = PatternFill(start_color=TEAL, end_color=TEAL, fill_type='solid')
        
        # Date row
        ws.row_dimensions[4].height = 22
        ws.merge_cells('A4:E4')
        ws['A4'] = report_date
        ws['A4'].font = Font(name='Calibri', size=10, color='666666')
        ws['A4'].fill = PatternFill(start_color=SNOW, end_color=SNOW, fill_type='solid')
        ws['A4'].alignment = Alignment(vertical='center', horizontal='center')
        for c in range(1, 6):
            ws.cell(row=4, column=c).fill = PatternFill(start_color=SNOW, end_color=SNOW, fill_type='solid')
        
        ws.row_dimensions[5].height = 8
        
        # === SHIFT HEADER ===
        row = 6
        shift_time = "08:00 - 18:00" if shift_type == "Day Shift" else "20:15 - 04:15"
        ws.merge_cells(f'A{row}:E{row}')
        ws[f'A{row}'] = f"{shift_type.upper()}  |  {shift_time}"
        ws[f'A{row}'].font = Font(name='Calibri', size=12, bold=True, color=WHITE)
        ws[f'A{row}'].fill = PatternFill(start_color=SQUID_INK, end_color=SQUID_INK, fill_type='solid')
        ws[f'A{row}'].alignment = Alignment(vertical='center', horizontal='center')
        for c in range(1, 6):
            ws.cell(row=row, column=c).fill = PatternFill(start_color=SQUID_INK, end_color=SQUID_INK, fill_type='solid')
        ws.row_dimensions[row].height = 25
        row += 1
        
        # Metrics row
        ws.row_dimensions[row].height = 24
        total_exceptions = len(excess) + len(less)
        ws[f'A{row}'] = f"Exceptions: {total_exceptions}"
        ws[f'A{row}'].font = Font(name='Calibri', size=9, bold=True, color=SQUID_INK)
        ws[f'A{row}'].fill = PatternFill(start_color=ICE_BLUE, end_color=ICE_BLUE, fill_type='solid')
        ws[f'B{row}'] = f"Excess: {len(excess)}  |  Less: {len(less)}"
        ws[f'B{row}'].font = Font(name='Calibri', size=9, color='666666')
        ws[f'B{row}'].fill = PatternFill(start_color=ICE_BLUE, end_color=ICE_BLUE, fill_type='solid')
        for c in range(3, 6):
            ws.cell(row=row, column=c).fill = PatternFill(start_color=ICE_BLUE, end_color=ICE_BLUE, fill_type='solid')
        row += 1
        
        # Spacer
        ws.row_dimensions[row].height = 6
        row += 1
        
        # === EXCESS TABLE ===
        ws.merge_cells(f'A{row}:E{row}')
        ws[f'A{row}'] = "EXCESS BREAK  ≥65 min"
        ws[f'A{row}'].font = Font(name='Calibri', size=10, bold=True, color=WHITE)
        ws[f'A{row}'].fill = PatternFill(start_color=CORAL, end_color=CORAL, fill_type='solid')
        ws[f'A{row}'].alignment = Alignment(vertical='center')
        for c in range(1, 6):
            ws.cell(row=row, column=c).fill = PatternFill(start_color=CORAL, end_color=CORAL, fill_type='solid')
        ws.row_dimensions[row].height = 20
        row += 1
        
        # Column headers
        headers = ['Employee ID', 'Name', 'Department', 'Mins', 'Repeat']
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col_idx, value=header)
            cell.font = Font(name='Calibri', size=8, bold=True, color=SQUID_INK)
            cell.fill = PatternFill(start_color=LIGHT_CORAL, end_color=LIGHT_CORAL, fill_type='solid')
        ws.row_dimensions[row].height = 16
        row += 1
        
        if len(excess) == 0:
            ws[f'A{row}'] = "No exceptions"
            ws[f'A{row}'].font = Font(name='Calibri', size=9, italic=True, color='AAAAAA')
            row += 1
        else:
            for i, (_, r) in enumerate(excess.iterrows()):
                ws.row_dimensions[row].height = 17
                is_repeat = r['Repeat'] != ""
                
                if is_repeat:
                    for c in range(1, 6):
                        ws.cell(row=row, column=c).fill = PatternFill(start_color=REPEAT_BG, end_color=REPEAT_BG, fill_type='solid')
                elif i % 2 == 0:
                    for c in range(1, 6):
                        ws.cell(row=row, column=c).fill = PatternFill(start_color=LIGHT_CORAL, end_color=LIGHT_CORAL, fill_type='solid')
                
                ws.cell(row=row, column=1, value=r['Employee ID']).font = Font(name='Calibri', size=9, color=DARK_TEXT)
                ws.cell(row=row, column=2, value=r['Employee Name']).font = Font(name='Calibri', size=9, color=DARK_TEXT)
                ws.cell(row=row, column=3, value=r['Department']).font = Font(name='Calibri', size=9, color=DARK_TEXT)
                ws.cell(row=row, column=4, value=r['Break (min)']).font = Font(name='Calibri', size=9, bold=True, color=CORAL)
                
                if is_repeat:
                    ws.cell(row=row, column=5, value=r['Repeat']).font = Font(name='Calibri', size=9, bold=True, color=REPEAT_RED)
                else:
                    ws.cell(row=row, column=5, value="").font = Font(name='Calibri', size=9)
                
                row += 1
        
        # Spacer
        ws.row_dimensions[row].height = 6
        row += 1
        
        # === LESS TABLE ===
        ws.merge_cells(f'A{row}:E{row}')
        ws[f'A{row}'] = "LESS BREAK  ≤55 min"
        ws[f'A{row}'].font = Font(name='Calibri', size=10, bold=True, color=WHITE)
        ws[f'A{row}'].fill = PatternFill(start_color=SUNSET, end_color=SUNSET, fill_type='solid')
        ws[f'A{row}'].alignment = Alignment(vertical='center')
        for c in range(1, 6):
            ws.cell(row=row, column=c).fill = PatternFill(start_color=SUNSET, end_color=SUNSET, fill_type='solid')
        ws.row_dimensions[row].height = 20
        row += 1
        
        # Column headers
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col_idx, value=header)
            cell.font = Font(name='Calibri', size=8, bold=True, color=SQUID_INK)
            cell.fill = PatternFill(start_color=LIGHT_SUNSET, end_color=LIGHT_SUNSET, fill_type='solid')
        ws.row_dimensions[row].height = 16
        row += 1
        
        if len(less) == 0:
            ws[f'A{row}'] = "No exceptions"
            ws[f'A{row}'].font = Font(name='Calibri', size=9, italic=True, color='AAAAAA')
            row += 1
        else:
            for i, (_, r) in enumerate(less.iterrows()):
                ws.row_dimensions[row].height = 17
                is_repeat = r['Repeat'] != ""
                
                if is_repeat:
                    for c in range(1, 6):
                        ws.cell(row=row, column=c).fill = PatternFill(start_color=REPEAT_BG, end_color=REPEAT_BG, fill_type='solid')
                elif i % 2 == 0:
                    for c in range(1, 6):
                        ws.cell(row=row, column=c).fill = PatternFill(start_color=LIGHT_SUNSET, end_color=LIGHT_SUNSET, fill_type='solid')
                
                ws.cell(row=row, column=1, value=r['Employee ID']).font = Font(name='Calibri', size=9, color=DARK_TEXT)
                ws.cell(row=row, column=2, value=r['Employee Name']).font = Font(name='Calibri', size=9, color=DARK_TEXT)
                ws.cell(row=row, column=3, value=r['Department']).font = Font(name='Calibri', size=9, color=DARK_TEXT)
                ws.cell(row=row, column=4, value=r['Break (min)']).font = Font(name='Calibri', size=9, bold=True, color=SUNSET)
                
                if is_repeat:
                    ws.cell(row=row, column=5, value=r['Repeat']).font = Font(name='Calibri', size=9, bold=True, color=REPEAT_RED)
                else:
                    ws.cell(row=row, column=5, value="").font = Font(name='Calibri', size=9)
                
                row += 1
        
        # Footer line
        for c in range(1, 6):
            ws.cell(row=row, column=c).border = Border(top=Side(style='medium', color=TEAL))
        
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer
    
    # --- Downloads ---
    st.markdown("---")
    st.markdown("### 📥 Downloads")
    
    col1, col2 = st.columns(2)
    
    with col1:
        excel_buffer = generate_excel(excess_df, less_df, report_date, shift_type)
        st.download_button(
            label="📊 Excel Report",
            data=excel_buffer,
            file_name=f"Break_Report_{short_date.replace('/', '-')}_{shift_type.replace(' ', '_')}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.document"
        )
    
    with col2:
        history_csv = updated_history.to_csv(index=False)
        st.download_button(
            label="📋 Updated History",
            data=history_csv,
            file_name="history.csv",
            mime="text/csv"
        )
    
    st.markdown("---")
    st.markdown("💡 **Download the updated history.csv each time and upload it next session!**")

# --- Criteria ---
st.markdown("---")
st.markdown("### ℹ️ How to use:")
st.markdown("1. (Optional) Upload previous history.csv for repeat tracking")
st.markdown("2. Upload your attendance CSV (supports both Dashboard export & FCLM raw data)")
st.markdown("3. View results + download report & updated history!")
st.markdown("### 📐 Criteria:")
st.markdown("- **Excess** = ≥65 min | **Less** = ≤55 min")
st.markdown("- Repeat offenders highlighted in red with count")

